from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from build_xlsx_portable import write_workbook
from image_to_rows import extract, merge_adjacent_empty_cells, typed_value


ROOT = Path(__file__).parent


class ImageToRowsRegressionTests(unittest.TestCase):
    def test_table4_keeps_all_colored_sections_and_footer_text(self) -> None:
        data = extract(ROOT / "example_pics" / "表4期权限仓20260609.png")

        expected_labels = [
            "大连商品交易所",
            "郑州商品交易所",
            "上海期货交易所",
            "上海国际能源交易中心",
            "中国金融期货交易所",
            "广州期货交易所",
        ]
        self.assertEqual([section["label"] for section in data["sections"]], expected_labels)
        self.assertEqual(
            [(len(section["cells"]), len(section["x_edges"]) - 1) for section in data["sections"]],
            [(11, 2), (17, 3), (22, 3), (7, 4), (4, 3), (4, 2)],
        )
        empty_merges = [
            merged
            for section in data["sections"]
            for merged in section["merged_cells"]
            if not str(merged.get("value", "")).strip()
        ]
        self.assertTrue(all(merged["r0"] != merged["r1"] or merged["c0"] != merged["c1"] for merged in empty_merges))
        notes = "\n".join(note["text"] for note in data.get("colored_notes", []))
        self.assertIn("分别计算", notes)
        footer = "\n".join(note["text"] for note in data.get("footer_notes", []))
        self.assertIn("400-700-7878", footer)

    def test_same_table6_template_keeps_the_same_column_structure(self) -> None:
        signatures = []
        for name in ("表6异常交易20260415.png", "表6异常交易20260714.png"):
            data = extract(ROOT / "example_pics" / name)
            signatures.append([len(section["x_edges"]) - 1 for section in data["sections"]])
            self.assertIn("套利、套保是否豁免", data["sections"][0]["cells"][0][4].replace("\n", ""))
            self.assertIn("套期保值", data["sections"][0]["cells"][1][4])
            self.assertIn("是", data["sections"][0]["cells"][1][4])

        self.assertEqual(signatures[0], signatures[1])

    def test_table6_keeps_the_main_block_and_does_not_add_fake_purple_headers(self) -> None:
        data = extract(ROOT / "example_pics" / "表6异常交易20260415.png")

        self.assertEqual(len(data["sections"]), 2)
        self.assertEqual([len(section["x_edges"]) - 1 for section in data["sections"]], [7, 10])
        empty_merges = [
            merged
            for section in data["sections"]
            for merged in section["merged_cells"]
            if not str(merged.get("value", "")).strip()
        ]
        self.assertTrue(all(merged["r0"] != merged["r1"] or merged["c0"] != merged["c1"] for merged in empty_merges))

        with TemporaryDirectory() as directory:
            output = Path(directory) / "table6.xlsx"
            write_workbook(data, output)
            sheet = load_workbook(output, data_only=False)["识别结果"]
            purple_rows = [
                row
                for row in range(1, sheet.max_row + 1)
                if sheet.cell(row, 1).fill.fgColor.rgb == "009524B2"
            ]
            self.assertEqual(purple_rows, [])

    def test_adjacent_empty_cells_are_merged_into_a_rectangle(self) -> None:
        cells = [["标题", "", ""], ["值", "", ""]]

        merged = merge_adjacent_empty_cells(cells, [])

        self.assertIn(
            {"r0": 0, "r1": 1, "c0": 1, "c1": 2, "value": "", "ocr_item_count": 0, "ocr_min_score": None, "ocr_max_score": None},
            merged,
        )

    def test_ocr_scalar_values_remain_exact_text(self) -> None:
        self.assertEqual(typed_value("12%"), "12%")
        self.assertEqual(typed_value("1,234"), "1,234")
        self.assertIsInstance(typed_value("0.125"), str)


if __name__ == "__main__":
    unittest.main()
