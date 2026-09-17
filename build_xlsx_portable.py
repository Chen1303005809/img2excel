#!/usr/bin/env python3
"""Portable XLSX writer for the image-to-rows JSON contract.

This file is intended for an ordinary Linux or Windows deployment where the
Codex-only artifact-tool runtime is not available.  It contains no OCR or
table-layout logic, so the recognition engine can evolve independently.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


PURPLE = "9524B2"
DARK_PURPLE = "4B0F72"
LIGHT_BORDER = "B7B7B7"
LIGHT_WARNING = "FFF2CC"
WARNING_TEXT = "6B4F00"

THIN = Side(style="thin", color=LIGHT_BORDER)
GRID_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def section_column_count(section: dict[str, Any]) -> int:
    rows = section.get("cells", [])
    return max(
        len(section.get("x_edges", [])) - 1,
        *(len(row) for row in rows),
        1,
    )


def text_length(value: Any) -> int:
    return len(value) if isinstance(value, str) else 0


def number_format(value: Any) -> str | None:
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return None
    if value != 0 and abs(value) < 1:
        return "0.0%"
    if isinstance(value, int) or float(value).is_integer():
        return "#,##0"
    return "#,##0.##"


def style_grid(ws, start_row: int, end_row: int, start_col: int, end_col: int, header: bool = False) -> None:
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=start_col, max_col=end_col):
        for cell in row:
            cell.border = GRID_BORDER
            cell.alignment = Alignment(
                horizontal="center" if header else "left",
                vertical="center",
                wrap_text=True,
            )
            cell.font = Font(name="Arial", size=9, bold=header, color="FFFFFF" if header else "222222")
            if header:
                cell.fill = PatternFill("solid", fgColor=DARK_PURPLE)


def set_row_height(ws, row_index: int, values: list[Any], end_col: int) -> None:
    longest = max((text_length(value) for value in values), default=0)
    line_count = max((str(value or "").count("\n") + 1 for value in values), default=1)
    ws.row_dimensions[row_index].height = 54 if longest > 90 or line_count >= 4 else 38 if longest > 42 or line_count >= 2 else 22


def write_workbook(data: dict[str, Any], output_path: Path) -> None:
    workbook = Workbook()
    result = workbook.active
    result.title = "识别结果"
    quality = workbook.create_sheet("识别质量")
    raw = workbook.create_sheet("OCR坐标")
    for sheet in (result, quality, raw):
        sheet.sheet_view.showGridLines = False

    sections = data.get("sections", [])
    max_columns = max((section_column_count(section) for section in sections), default=1)
    last_column = get_column_letter(max_columns)

    result.merge_cells(f"A2:{last_column}2")
    result["A2"] = data.get("title") or "图片表格识别结果"
    result["A2"].font = Font(name="Arial", size=14, bold=True, color=DARK_PURPLE)
    result.row_dimensions[2].height = 26

    source_row = ["" for _ in range(max_columns)]
    source_row[0] = "来源：用户提供图片"
    if max_columns > 2:
        source_row[2] = "分辨率"
    if max_columns > 3:
        source = data.get("source", {})
        source_row[3] = f"{source.get('width', '')} × {source.get('height', '')}"
    for column, value in enumerate(source_row, start=1):
        result.cell(3, column, value)
        result.cell(3, column).font = Font(name="Arial", size=9, italic=True, color="666666")

    result.merge_cells(f"A4:{last_column}4")
    result["A4"] = "说明：按图像自身的网格线和文字布局自动推断表格。带换行文字保留在同一单元格，OCR 坐标与分数见 OCR坐标。"
    result["A4"].font = Font(name="Arial", size=9, italic=True, color="666666")

    widths = [18, 22, 25, 38, 30] + [18] * max(0, max_columns - 5)
    for column, width in enumerate(widths[:max_columns], start=1):
        result.column_dimensions[get_column_letter(column)].width = width

    current_row = 5
    for section in sections:
        section_columns = section_column_count(section)
        label = str(section.get("label", "")).strip()
        if label:
            section_last = get_column_letter(section_columns)
            result.merge_cells(f"A{current_row}:{section_last}{current_row}")
            result.cell(current_row, 1, label)
            result.cell(current_row, 1).fill = PatternFill("solid", fgColor=PURPLE)
            result.cell(current_row, 1).font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            result.cell(current_row, 1).alignment = Alignment(horizontal="center", vertical="center")
            result.row_dimensions[current_row].height = 21
            current_row += 1

        rows = section.get("cells", [])
        for row_offset, row_values in enumerate(rows):
            row_index = current_row + row_offset
            padded = list(row_values) + [""] * max(0, section_columns - len(row_values))
            for column, value in enumerate(padded, start=1):
                cell = result.cell(row_index, column, value)
                cell.number_format = number_format(value) or "General"
            style_grid(result, row_index, row_index, 1, section_columns)
            set_row_height(result, row_index, padded, section_columns)

        for merged in section.get("merged_cells", []):
            if merged["r0"] == merged["r1"] and merged["c0"] == merged["c1"]:
                continue
            row0 = current_row + merged["r0"]
            row1 = current_row + merged["r1"]
            col0 = merged["c0"] + 1
            col1 = merged["c1"] + 1
            result.merge_cells(start_row=row0, start_column=col0, end_row=row1, end_column=col1)
            result.cell(row0, col0).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        current_row += len(rows)

    for note in data.get("colored_notes", []):
        text = str(note.get("text", "")).strip()
        if not text:
            continue
        result.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_columns)
        cell = result.cell(current_row, 1, text)
        cell.fill = PatternFill("solid", fgColor=PURPLE)
        cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        result.row_dimensions[current_row].height = min(72, max(24, 20 * (text.count("\n") + 1)))
        current_row += 1

    for note in data.get("footer_notes", []):
        text = str(note.get("text", "")).strip()
        if not text:
            continue
        result.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_columns)
        cell = result.cell(current_row, 1, text)
        cell.font = Font(name="Arial", size=9, color="222222")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        result.row_dimensions[current_row].height = min(54, max(22, 20 * (text.count("\n") + 1)))
        current_row += 1

    # Quality view
    quality.merge_cells("A2:D2")
    quality["A2"] = "图片转 Excel 实验质量"
    quality["A2"].font = Font(name="Arial", size=14, bold=True, color=DARK_PURPLE)
    quality.row_dimensions[2].height = 26
    quality.append([])
    quality.append(["指标", "值", "说明"])
    style_grid(quality, 4, 4, 1, 3, header=True)
    source = data.get("source", {})
    metrics = data.get("metrics", {})
    strategy_counts = metrics.get("strategy_counts", {})
    metric_rows = [
        ["输入文件", source.get("filename", ""), "仅使用用户提供的附件"],
        ["输入分辨率", source.get("width", 0) * source.get("height", 0), f"{source.get('width', '')} × {source.get('height', '')} 像素"],
        ["检测到的表格块", metrics.get("section_count", 0), "按网格线和可选彩色标题带自动发现"],
        ["OCR 文字框", metrics.get("ocr_box_count", 0), "分区识别后合计"],
        ["OCR 平均置信度", metrics.get("ocr_avg_score"), "模型分数，不等同于字符准确率"],
        ["低于 0.75 的文字框", metrics.get("ocr_low_score_count", 0), "仍需人工抽查高分误识别"],
        ["二次修正", len(data.get("corrections", [])), "对疑似低置信度单元格做裁边复识别"],
        ["识别链路", "OpenCV + RapidOCR + ONNX Runtime CPU", "几何与文字识别解耦"],
        ["表格策略", "；".join(f"{key}: {value}" for key, value in strategy_counts.items()), "规则表优先，无边框时使用 OCR 排版后备"],
        ["彩色标题带", metrics.get("colored_band_count", 0), "优先定位分区、彩色说明框和页首表头；不足时回退到线段检测"],
        ["自动后备分支", "已使用" if metrics.get("fallback_used") else "未使用", "后备分支用于处理无边框或断线表格"],
    ]
    for row in metric_rows:
        quality.append(row)
    style_grid(quality, 5, 4 + len(metric_rows), 1, 3)
    quality["B6"].number_format = "#,##0"
    quality["B9"].number_format = "0.0%"
    section_header = 5 + len(metric_rows) + 1
    quality.cell(section_header, 1, "表格块")
    quality.cell(section_header, 2, "OCR框 / 平均分 / 低分框")
    quality.cell(section_header, 3, "网格行 × 列 / 合并块 / 修正 / 策略")
    style_grid(quality, section_header, section_header, 1, 3, header=True)
    for index, section in enumerate(sections, start=1):
        quality.append([
            f"{section.get('id', f'T{index:02d}')} {section.get('label', '')}",
            f"{section.get('ocr_count', 0)} / {section.get('ocr_avg_score')} / {section.get('ocr_low_score_count', 0)}",
            f"{len(section.get('cells', []))} × {section_column_count(section)} / {len(section.get('merged_cells', []))} / {sum(item.get('section') == index for item in data.get('corrections', []))} / {section.get('strategy', '')}",
        ])
    section_start = section_header + 1
    section_end = section_start + max(0, len(sections) - 1)
    if sections:
        style_grid(quality, section_start, section_end, 1, 3)
    conclusion = section_end + 2
    quality.merge_cells(start_row=conclusion, start_column=1, end_row=conclusion, end_column=4)
    quality.cell(conclusion, 1, "说明：默认按图像自身的线段和文字布局自动推断，不要求每张图手工调列线。低置信度、数字、百分号、窄列和长合并表头仍建议抽查。")
    quality.cell(conclusion, 1).fill = PatternFill("solid", fgColor=LIGHT_WARNING)
    quality.cell(conclusion, 1).font = Font(name="Arial", size=9, color=WARNING_TEXT)
    quality.cell(conclusion, 1).alignment = Alignment(vertical="center", wrap_text=True)
    quality.row_dimensions[conclusion].height = 48
    for column, width in {"A": 28, "B": 32, "C": 42, "D": 16}.items():
        quality.column_dimensions[column].width = width

    # Raw OCR evidence view
    headers = ["分区", "行", "列", "x", "y", "宽", "高", "中心x", "中心y", "文本", "分数"]
    raw.append(headers)
    for item in data.get("ocr_boxes", []):
        raw.append([
            item.get("section"), item.get("row"), item.get("column"), item.get("x"), item.get("y"),
            item.get("w"), item.get("h"), item.get("cx"), item.get("cy"), item.get("text"), item.get("score"),
        ])
    style_grid(raw, 1, max(1, raw.max_row), 1, len(headers), header=True)
    for column in range(4, 10):
        for row in range(2, raw.max_row + 1):
            raw.cell(row, column).number_format = "0.00"
    for row in range(2, raw.max_row + 1):
        raw.cell(row, 11).number_format = "0.0000"
    raw.freeze_panes = "A2"
    if raw.max_row > 1:
        table = Table(displayName="OCRBoxes", ref=f"A1:K{raw.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        raw.add_table(table)
    for column, width in {"A": 9, "B": 8, "C": 8, "D": 12, "E": 12, "F": 10, "G": 10, "H": 12, "I": 12, "J": 48, "K": 10}.items():
        raw.column_dimensions[column].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description="Write a portable XLSX from image_to_rows JSON.")
    parser.add_argument("json_path", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    write_workbook(json.loads(args.json_path.read_text(encoding="utf-8")), args.output)
    print(f"XLSX_SAVED {args.output}")


if __name__ == "__main__":
    main()
