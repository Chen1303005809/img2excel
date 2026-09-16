from __future__ import annotations

import hashlib
import json
from copy import deepcopy
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from backend.app.crawling import CrawlSnapshot, ImageCandidateData
from backend.app.models import Artifact, ImageCandidate, Run, Source, utc_now
from backend.app.pipeline import DownloadedImage, PipelineRunner

from .conftest import make_document


class FakeCrawler:
    def __init__(self, candidates: tuple[ImageCandidateData, ...]):
        self.candidates = candidates
        self.calls: list[str] = []

    async def crawl(self, url: str) -> CrawlSnapshot:
        self.calls.append(url)
        return CrawlSnapshot(
            requested_url=url,
            final_url=url + "#final",
            title="测试页面",
            success=True,
            status_code=200,
            error_message="",
            candidates=self.candidates,
            metadata={
                "requested_url": url,
                "final_url": url + "#final",
                "success": True,
                "status_code": 200,
                "content_images": [candidate.__dict__ for candidate in self.candidates],
            },
        )


class SequenceExtractor:
    def __init__(self, documents: list[dict]):
        self.documents = [deepcopy(document) for document in documents]
        self.paths: list[Path] = []

    def extract(self, image_path: Path) -> dict:
        self.paths.append(image_path)
        return deepcopy(self.documents.pop(0))


class FakeDownloader:
    async def download(self, candidate: ImageCandidateData) -> DownloadedImage:
        content = f"fake-image-{candidate.ordinal}".encode()
        return DownloadedImage(
            content=content,
            media_type="image/png",
            extension=".png",
            sha256=hashlib.sha256(content).hexdigest(),
        )


def _create_source_and_run(session_factory, *, source_id: str | None = None, baseline_run_id: str | None = None) -> tuple[Source, Run]:
    source_id = source_id or str(uuid4())
    run_id = str(uuid4())
    url = "https://example.com/table"
    with session_factory() as session:
        source = session.get(Source, source_id)
        if source is None:
            source = Source(
                id=source_id,
                name="测试来源",
                url=url,
                normalized_url=url,
                profile_key="yafco_image",
                enabled=True,
            )
            session.add(source)
            session.commit()
        run = Run(
            id=run_id,
            source_id=source_id,
            requested_url=url,
            normalized_url=url,
            profile_key="yafco_image",
            baseline_run_id=baseline_run_id,
            status="queued",
            stage="queued",
            progress=0,
            message="等待 Worker 处理",
        )
        session.add(run)
        session.commit()
        session.refresh(source)
        session.refresh(run)
        return source, run


def _runner(db_env, crawler: FakeCrawler, extractor: SequenceExtractor) -> PipelineRunner:
    settings, _, session_factory, store = db_env
    runner = PipelineRunner(settings, session_factory, store, crawler, extractor=extractor)
    runner.downloader = FakeDownloader()
    return runner


@pytest.mark.asyncio
async def test_pipeline_persists_original_outputs_and_comparison(db_env):
    candidates = (ImageCandidateData("/source.png", "https://example.com/source.png", 0, "表格", 800, 300),)
    crawler = FakeCrawler(candidates)
    extractor = SequenceExtractor([make_document("100"), make_document("200")])
    runner = _runner(db_env, crawler, extractor)
    _, first_run = _create_source_and_run(db_env[2])
    await runner.process(first_run.id)

    with db_env[2]() as session:
        persisted = session.get(Run, first_run.id)
        assert persisted is not None
        assert persisted.status == "succeeded"
        assert persisted.final_url.endswith("#final")
        first_artifacts = {item.kind: item for item in session.query(Artifact).filter_by(run_id=first_run.id)}
        assert {"crawl_metadata", "source_image", "recognized_json", "recognized_xlsx"} <= first_artifacts.keys()
        json_payload = json.loads(db_env[3].absolute_path(first_artifacts["recognized_json"].relative_path).read_text(encoding="utf-8"))
        assert json_payload["schema_version"] == 1
        assert json_payload["document"]["version"] == 2
        workbook = load_workbook(db_env[3].absolute_path(first_artifacts["recognized_xlsx"].relative_path), read_only=True)
        assert {"识别结果", "识别质量", "OCR坐标"} <= set(workbook.sheetnames)
        workbook.close()

    _, second_run = _create_source_and_run(db_env[2], source_id=first_run.source_id, baseline_run_id=first_run.id)
    await runner.process(second_run.id)
    with db_env[2]() as session:
        persisted = session.get(Run, second_run.id)
        assert persisted is not None
        assert persisted.status == "succeeded"
        assert persisted.comparison_summary["has_baseline"] is True
        assert persisted.comparison_summary["changed_cells"] == 1


@pytest.mark.asyncio
async def test_pipeline_requires_explicit_choice_for_multiple_images(db_env):
    candidates = (
        ImageCandidateData("/one.png", "https://example.com/one.png", 0, "第一张", 800, 300),
        ImageCandidateData("/two.png", "https://example.com/two.png", 1, "第二张", 800, 300),
    )
    crawler = FakeCrawler(candidates)
    extractor = SequenceExtractor([make_document("300")])
    runner = _runner(db_env, crawler, extractor)
    _, run = _create_source_and_run(db_env[2])
    await runner.process(run.id)

    with db_env[2]() as session:
        waiting = session.get(Run, run.id)
        assert waiting is not None
        assert waiting.status == "awaiting_image_selection"
        persisted_candidates = list(session.query(ImageCandidate).filter_by(run_id=run.id).order_by(ImageCandidate.ordinal))
        assert len(persisted_candidates) == 2
        assert not any(item.selected for item in persisted_candidates)
        persisted_candidates[1].selected = True
        waiting.status = "queued"
        waiting.stage = "queued"
        waiting.progress = 25
        session.commit()

    await runner.process(run.id)
    with db_env[2]() as session:
        persisted = session.get(Run, run.id)
        assert persisted is not None
        assert persisted.status == "succeeded"
        assert len(extractor.paths) == 1


@pytest.mark.asyncio
async def test_pipeline_fails_and_keeps_crawl_metadata_when_no_image(db_env):
    crawler = FakeCrawler(())
    runner = _runner(db_env, crawler, SequenceExtractor([make_document()]))
    _, run = _create_source_and_run(db_env[2])
    await runner.process(run.id)

    with db_env[2]() as session:
        persisted = session.get(Run, run.id)
        assert persisted is not None
        assert persisted.status == "failed"
        assert persisted.error_code == "no_content_image"
        assert session.query(Artifact).filter_by(run_id=run.id, kind="crawl_metadata").count() == 1
